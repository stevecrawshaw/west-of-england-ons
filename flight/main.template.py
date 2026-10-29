"""Refresh the ons_explore MotherDuck database from ONS Explore Local Statistics.

Do not edit flight/ons_explore_refresh/main.py directly: it is generated from this
template and the sql/ directory by scripts/build_flight.py.

Sources (fixed URLs, overwritten in place by ONS):
- all-datasets.csv   long-format observations for every indicator -> fact_ons_data
- all-datasets.xlsx  one sheet per indicator with description, unit and source
                     (absent from the CSV)                         -> dim_ons_data

Config (environment variables):
- FORCE=1         reload even if the ONS files have not changed
- LOCAL_DB        path to a local DuckDB file named ons_explore.duckdb (testing)
- SOURCE_DIR      directory holding all-datasets.csv and all-datasets.xlsx,
                  used instead of downloading (testing)
"""

import csv
import gzip
import io
import os
import re
import shutil
import sys
from pathlib import Path

import duckdb
import openpyxl
import requests

BASE_URL = "https://www.ons.gov.uk/explore-local-statistics/files"
CSV_URL = f"{BASE_URL}/all-datasets.csv"
XLSX_URL = f"{BASE_URL}/all-datasets.xlsx"
HEADERS = {"User-Agent": "Mozilla/5.0 (ons_explore_refresh flight)"}
WORK_DIR = Path("/tmp/ons_explore_refresh")

EXPECTED_COLUMNS = [
    "indicator", "areacd", "areanm", "period", "sex", "age",
    "value", "lci_95", "uci_95", "status",
]
# Refuse to replace the tables if the new load is much smaller than the current one
MIN_ROW_RATIO = 0.9
MIN_INDICATOR_RATIO = 0.9
ENGLISH_LA_PREFIXES = ("E06", "E07", "E08", "E09")

# --- SQL, inlined from sql/ by scripts/build_flight.py -----------------------
SQL_BASE_TABLES = r"""@@SQL_00@@"""
SQL_DIM_VIZ = r"""@@SQL_01@@"""
SQL_EXPORT_VIEWS = r"""@@SQL_02@@"""
INDICATOR_CATEGORIES_CSV = r"""@@CATEGORIES@@"""
# -----------------------------------------------------------------------------


def log(message: str) -> None:
    print(message, flush=True)


def split_statements(script: str) -> list[str]:
    """Split a SQL script on semicolons, ignoring those in strings and comments."""
    statements: list[str] = []
    buf: list[str] = []
    in_string = False
    i = 0
    while i < len(script):
        ch = script[i]
        if in_string:
            buf.append(ch)
            if ch == "'":
                if script[i + 1 : i + 2] == "'":  # escaped quote
                    buf.append("'")
                    i += 1
                else:
                    in_string = False
        elif ch == "'":
            in_string = True
            buf.append(ch)
        elif script.startswith("--", i):
            while i < len(script) and script[i] != "\n":
                i += 1
            continue
        elif ch == ";":
            statement = "".join(buf).strip()
            if statement:
                statements.append(statement)
            buf = []
        else:
            buf.append(ch)
        i += 1
    tail = "".join(buf).strip()
    if tail:
        statements.append(tail)
    return statements


def run_script(con: duckdb.DuckDBPyConnection, script: str) -> None:
    for statement in split_statements(script):
        con.execute(statement)


def connect() -> duckdb.DuckDBPyConnection:
    local_db = os.environ.get("LOCAL_DB")
    if local_db:
        return duckdb.connect(local_db)
    con = duckdb.connect("md:")
    con.execute("CREATE DATABASE IF NOT EXISTS ons_explore")
    con.close()
    return duckdb.connect("md:ons_explore?attach_mode=single")


def remote_stamp(url: str) -> str:
    """Last-Modified (falling back to ETag) of a source file; '' if neither is sent."""
    response = requests.head(url, headers=HEADERS, timeout=60, allow_redirects=True)
    response.raise_for_status()
    return response.headers.get("Last-Modified") or response.headers.get("ETag") or ""


def download(url: str, dest: Path) -> None:
    log(f"Downloading {url}")
    with requests.get(url, headers=HEADERS, stream=True, timeout=(30, 300)) as response:
        response.raise_for_status()
        with dest.open("wb") as f:
            shutil.copyfileobj(response.raw, f, length=1 << 20)
    # ONS labels the CSV application/gzip; guard against it ever being gzipped.
    with dest.open("rb") as f:
        magic = f.read(2)
    if magic == b"\x1f\x8b":
        plain = dest.with_suffix(dest.suffix + ".plain")
        with gzip.open(dest, "rb") as src, plain.open("wb") as out:
            shutil.copyfileobj(src, out, length=1 << 20)
        plain.replace(dest)
    log(f"  {dest.stat().st_size / 1e6:.1f} MB")


def fetch_sources() -> tuple[Path, Path]:
    source_dir = os.environ.get("SOURCE_DIR")
    if source_dir:
        return Path(source_dir) / "all-datasets.csv", Path(source_dir) / "all-datasets.xlsx"
    shutil.rmtree(WORK_DIR, ignore_errors=True)
    WORK_DIR.mkdir(parents=True)
    csv_path, xlsx_path = WORK_DIR / "all-datasets.csv", WORK_DIR / "all-datasets.xlsx"
    download(CSV_URL, csv_path)
    download(XLSX_URL, xlsx_path)
    return csv_path, xlsx_path


def read_indicator_metadata(xlsx_path: Path) -> list[dict]:
    """Parse title, description, source and unit from each numbered indicator sheet.

    Each sheet starts with: A1 "<Indicator> [note n]", A2 description, a
    "Source: ..." row, then a header row beginning "Area code" whose period
    columns end in the unit, e.g. "2023-06-30 (Persons)".
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    rows = []
    for name in wb.sheetnames:
        if not re.fullmatch(r"\d+", name):
            continue
        head = list(wb[name].iter_rows(min_row=1, max_row=10, values_only=True))
        first_col = [r[0] if r else None for r in head]
        indicator = re.sub(r"\s*\[note\s*\d+\]$", "", str(first_col[0]))
        sources = [str(c) for c in first_col if isinstance(c, str) and c.startswith("Source")]
        unit = None
        for i, cell in enumerate(first_col):
            if cell == "Area code":
                for header in head[i]:
                    match = re.search(r"\(([^)]+)\)$", str(header or ""))
                    if match:
                        unit = match.group(1)
                        break
                break
        rows.append(
            {
                "sheet": int(name),
                "indicator": indicator,
                "description": first_col[1],
                "unit": unit,
                "source": " | ".join(sources),
            }
        )
    wb.close()
    return rows


def validate_csv(con: duckdb.DuckDBPyConnection, csv_path: Path) -> tuple[int, int]:
    with csv_path.open(encoding="utf-8", newline="") as f:
        header = next(csv.reader(f))
    if header != EXPECTED_COLUMNS:
        sys.exit(f"Unexpected CSV columns: {header}")
    rows, indicators = con.execute(
        "SELECT count(*), count(DISTINCT indicator) FROM stage_fact"
    ).fetchone()
    return rows, indicators


def main() -> None:
    force = os.environ.get("FORCE") == "1"
    con = connect()
    run_script(con, SQL_BASE_TABLES)

    if not os.environ.get("SOURCE_DIR"):
        csv_stamp, xlsx_stamp = remote_stamp(CSV_URL), remote_stamp(XLSX_URL)
        last = con.execute(
            "SELECT csv_last_modified, xlsx_last_modified FROM ons_explore.main.etl_load_log "
            "ORDER BY loaded_at DESC LIMIT 1"
        ).fetchone()
        if not force and last == (csv_stamp, xlsx_stamp) and csv_stamp and xlsx_stamp:
            log(f"ONS files unchanged since last load ({csv_stamp}); nothing to do.")
            return
    else:
        csv_stamp = xlsx_stamp = "local"

    csv_path, xlsx_path = fetch_sources()

    # Stage and validate before touching the live tables.
    con.execute(
        f"""
        CREATE OR REPLACE TEMP TABLE stage_fact AS
        SELECT * FROM read_csv(
            '{csv_path}', header = true,
            columns = {{
                'indicator': 'VARCHAR', 'areacd': 'VARCHAR', 'areanm': 'VARCHAR',
                'period': 'DATE', 'sex': 'VARCHAR', 'age': 'VARCHAR',
                'value': 'DOUBLE', 'lci_95': 'VARCHAR', 'uci_95': 'VARCHAR',
                'status': 'VARCHAR'
            }}
        )
        """
    )
    new_rows, new_indicators = validate_csv(con, csv_path)
    old_rows, old_indicators = con.execute(
        "SELECT count(*), count(DISTINCT indicator) FROM ons_explore.main.fact_ons_data"
    ).fetchone()
    log(f"Staged {new_rows:,} rows / {new_indicators} indicators "
        f"(current: {old_rows:,} / {old_indicators})")
    if old_rows and (
        new_rows < MIN_ROW_RATIO * old_rows
        or new_indicators < MIN_INDICATOR_RATIO * old_indicators
    ):
        sys.exit("New load is much smaller than the current tables; aborting.")

    # Dimension: metadata for indicators with English local authority data.
    la_filter = " OR ".join(f"areacd LIKE '{p}%'" for p in ENGLISH_LA_PREFIXES)
    english = {
        r[0]
        for r in con.execute(
            f"SELECT DISTINCT indicator FROM stage_fact WHERE {la_filter}"
        ).fetchall()
    }
    meta = [m for m in read_indicator_metadata(xlsx_path) if m["indicator"] in english]
    categories = dict(csv.reader(io.StringIO(INDICATOR_CATEGORIES_CSV)))
    for m in meta:
        m["category"] = categories.get(m["indicator"])
    missing_meta = english - {m["indicator"] for m in meta}
    if missing_meta:
        sys.exit(f"No metadata sheet for indicators with LA data: {sorted(missing_meta)}")
    unmapped = sorted(m["indicator"] for m in meta if m["category"] is None)
    if unmapped:
        log(f"WARNING: no category for {unmapped}; add them to sql/indicator_categories.csv "
            "and rebuild the flight. They are loaded but will not appear in category cards.")

    dim_csv = (WORK_DIR if not os.environ.get("SOURCE_DIR") else Path("/tmp")) / "dim_stage.csv"
    dim_csv.parent.mkdir(parents=True, exist_ok=True)
    with dim_csv.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f, ["sheet", "indicator", "description", "unit", "source", "category"]
        )
        writer.writeheader()
        writer.writerows(meta)

    # Swap data in one transaction; views are dropped first because DuckDB
    # blocks ALTER/COMMENT on tables that views depend on.
    con.execute("DROP VIEW IF EXISTS ons_explore.main.export_ons_values")
    con.execute("DROP VIEW IF EXISTS ons_explore.main.export_ons_indicators")
    con.execute("BEGIN")
    try:
        con.execute("DELETE FROM ons_explore.main.fact_ons_data")
        con.execute("INSERT INTO ons_explore.main.fact_ons_data SELECT * FROM stage_fact")
        con.execute("DELETE FROM ons_explore.main.dim_ons_data")
        con.execute(
            f"""
            INSERT INTO ons_explore.main.dim_ons_data (sheet, indicator, description, unit, source, category)
            SELECT sheet, indicator, description, unit, source, category
            FROM read_csv('{dim_csv}', header = true, all_varchar = false,
                          types = {{'sheet': 'BIGINT'}})
            """
        )
        con.execute("COMMIT")
    except Exception:
        con.execute("ROLLBACK")
        raise

    run_script(con, SQL_DIM_VIZ)
    run_script(con, SQL_EXPORT_VIEWS)

    # Every category must still have its default indicator.
    defaults = con.execute(
        "SELECT count(*) FILTER (WHERE is_category_default), "
        "count(DISTINCT category) FILTER (WHERE category IS NOT NULL) "
        "FROM ons_explore.main.dim_ons_data"
    ).fetchone()
    if defaults[0] != defaults[1]:
        sys.exit(
            f"{defaults[0]} default indicators for {defaults[1]} categories: "
            "an indicator named in sql/01_dim_viz_metadata.sql has been renamed or removed."
        )

    con.execute(
        "INSERT INTO ons_explore.main.etl_load_log VALUES "
        "(current_timestamp::TIMESTAMP, ?, ?, ?, ?, ?, ?)",
        [csv_stamp, xlsx_stamp, new_rows, new_indicators, len(meta), ", ".join(unmapped)],
    )
    log(f"Loaded {new_rows:,} fact rows, {len(meta)} dim indicators.")
    shutil.rmtree(WORK_DIR, ignore_errors=True)


if __name__ == "__main__":
    main()
